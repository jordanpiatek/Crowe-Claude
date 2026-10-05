from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

wb = Workbook()
HF = PatternFill("solid", fgColor="002060"); HFont = Font(bold=True, color="FFFFFF", name="Arial", size=10)
thin = Side(style="thin", color="BFBFBF"); BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
BODY = Font(name="Arial", size=9)

def sheet(ws, headers, rows, widths, freeze="A2"):
    ws.append(headers)
    for c in ws[1]:
        c.fill = HF; c.font = HFont; c.alignment = Alignment(wrap_text=True, vertical="center"); c.border = BORDER
    for r in rows: ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BODY; c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = BORDER
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze; ws.row_dimensions[1].height = 32
    ws.auto_filter.ref = ws.dimensions

# ---- README
ws = wb.active; ws.title = "README"
lines = [
 ("Ezurio – Requirements & Observations Tracker", True),
 ("Oct 19–21, 2026 onsite discovery (Gateworks SLO + Irvine). INTERNAL working file.", False),
 ("", False),
 ("Tabs", True),
 ("1 Requirements – candidate requirements seeded from the engagement letter 'Key Areas of Focus'. Status = Hypothesis until validated onsite. Add rows as you go; keep IDs sequential.", False),
 ("2 Observations – raw observation log by session (what we saw/heard). Link each to a requirement ID where relevant.", False),
 ("3 Irvine Backlog – optimization ideas for Irvine with value/effort; cross-check against the Azure DevOps backlog to avoid duplicates.", False),
 ("4 GWTrack Features – feature inventory captured during the GWTrack demo (what exists, where it should live: GWTrack / NetSuite / Arena / MES).", False),
 ("5 Option Scoring – weighted scoring of system Options A (GWTrack), B (Manufacturing Mobile / Advanced Mfg), C (3rd-party MES / SuiteApp). Footprint options are NOT scored onsite.", False),
 ("6 Open Items – follow-ups, documents requested and owners.", False),
 ("", False),
 ("Conventions", True),
 ("Priority: Must / Should / Could (MoSCoW). Fit: Native / Config / Custom / Gap / TBD.", False),
 ("ITAR/CUI flag: Y if the capability/data is potentially controlled – do not share the underlying detail with non-US persons until compliance confirms.", False),
 ("Confidentiality: no references to acquisition activity, a third site or consolidation anywhere in this file.", False),
]
for t, b in lines:
    ws.append([t]); ws.cell(ws.max_row, 1).font = Font(name="Arial", size=12 if b else 10, bold=b)
ws.column_dimensions["A"].width = 130
for r in ws.iter_rows():
    for c in r: c.alignment = Alignment(wrap_text=True, vertical="top")

# ---- Requirements
ws = wb.create_sheet("1 Requirements")
H = ["Req ID", "Area", "Requirement (what the future system must do)", "Source", "Site", "Session", "Current state: GWTrack", "Current state: Irvine (NetSuite/MM)",
     "Priority", "Fit A: GWTrack", "Fit B: NetSuite MM/Adv Mfg", "Fit C: 3rd-party MES/SuiteApp", "ITAR/CUI flag", "Owner", "Status", "Notes / evidence"]
R = [
 ("Work execution", "Release work orders to the floor with routing, work instructions and revision control", "EL focus: GWTrack fit", "Both", "1.3 / 3.2"),
 ("Work execution", "Operators start / pause / complete operations at the station with scan-driven transactions", "EL focus: MM timekeeping", "Both", "1.3 / 3.3"),
 ("Work execution", "Enforce operation sequence and validations (no out-of-sequence or missing-prior-step completions)", "EL focus: process control gaps", "Both", "1.3"),
 ("Work execution", "Handle partial quantities and rework loops per operation", "MM process outline", "Irvine", "3.3"),
 ("Material control", "Prevent / warn on starting a run without sufficient committed component material", "EL focus: material availability checks", "Irvine", "3.3"),
 ("Material control", "Support kitting / staging with one-WO-per-location control and visibility of under-picked items", "Irvine WMS process", "Irvine", "3.2"),
 ("Material control", "Issue / backflush components with accurate inventory detail (bin / lot)", "Irvine WMS process", "Both", "1.3 / 3.2"),
 ("Material control", "Manage substitutions / alternates and shortages with expedite visibility", "Ops walk", "Both", "1.2"),
 ("Costing", "Block or flag production when an assembly has no standard cost; maintain standard cost with approvals", "EL focus: standard cost enforcement", "Irvine", "3.3"),
 ("Costing", "Post WIP, scrap, labor and variances to NetSuite accurately and on time", "Ops walk", "Both", "1.4 / 3.3"),
 ("Labor / timekeeping", "Capture labor by operation/job (badge-in, pause, end) including indirect, rework and multi-operator time", "EL focus: timekeeping gaps", "Both", "1.4 / 3.3"),
 ("Labor / timekeeping", "Provide labor efficiency and cost reporting; hand-off to payroll if required", "EL focus: timekeeping gaps", "Both", "1.4 / 3.3"),
 ("Traceability", "Maintain lot / serial / as-built genealogy and answer containment queries quickly", "GWTrack features", "Gateworks", "1.3"),
 ("Traceability", "Retain test results and operator / equipment records per unit", "GWTrack features", "Gateworks", "1.3"),
 ("Quality", "Capture defects at point of work with coded reasons and produce defect heatmaps", "EL focus: quality tracking", "Both", "1.4 / 3.4"),
 ("Quality", "Manage NCR / rework / scrap disposition integrated to inventory and cost", "EL focus: quality tracking", "Both", "1.4 / 3.4"),
 ("Quality", "Track customer returns (RMA) from receipt through disposition with trace to original build", "EL focus: manual RMA", "Both", "1.4 / 3.4"),
 ("Quality", "Support corrective action (CAPA / 8D) and audit-ready quality records", "Ops walk", "Gateworks", "1.4"),
 ("Inventory / WMS", "Post-production reconciliation of issued vs. actual components without manual template/import steps", "Irvine dekit process", "Irvine", "3.4"),
 ("Inventory / WMS", "Prevent partial-pick side effects (task closes / reverse picks) or handle them gracefully", "Irvine WMS process", "Irvine", "3.2"),
 ("Engineering / PLM", "Receive BOM / routing / ECO changes from Arena (or successor) without re-keying", "Client update: Arena rationalization", "Gateworks", "2.1"),
 ("Integration", "Define system of record by function across GWTrack / NetSuite / Arena and eliminate duplicate masters", "Client update", "Gateworks", "2.1"),
 ("Integration", "Integrate with floor equipment, test stations and label printers", "GWTrack features", "Gateworks", "1.3 / 2.1"),
 ("Compliance", "Support ITAR / CUI handling: access by US-person, data residency and segregation", "EL focus: ITAR / CUI", "Gateworks", "2.2"),
 ("Compliance", "Provide electronic records, audit trail and retention to meet customer / quality-system audits", "Ops walk", "Both", "2.2"),
 ("Security", "Role-based access and segregation of duties for floor and office users", "Ops walk", "Both", "2.2"),
 ("Reporting", "Provide real-time floor dashboards and daily operational reports (WIP, queue, shortages, yield)", "GWTrack features", "Both", "1.3 / 2.1"),
 ("Reporting", "Ad hoc / saved reporting without developer support", "Ops walk", "Both", "2.1 / 3.4"),
 ("Performance", "Station transactions respond within agreed time at peak load", "Client update: keep performance", "Both", "2.3"),
 ("Performance", "Continue operating (or recover cleanly) when network / ERP is unavailable", "Client update: keep automation", "Gateworks", "2.3"),
 ("Automation", "Preserve existing GWTrack automations (auto-transactions, auto-completes, auto-labeling) – list to be confirmed", "Client update: keep automation", "Gateworks", "1.3 / 2.3"),
 ("Scalability", "Scale to additional users / lines / entities without redesign", "EL / assessment: GWTrack not built to extend", "Both", "2.3"),
 ("Supportability", "Vendor-supported platform with clear ownership; remove key-person risk", "Ops walk", "Gateworks", "2.1"),
]
rows = []
for i, (a, req, src, site, ses) in enumerate(R, 1):
    rows.append((f"REQ-{i:03d}", a, req, src, site, ses, "", "", "", "TBD", "TBD", "TBD", "", "", "Hypothesis", ""))
sheet(ws, H, rows, [9, 16, 58, 28, 10, 10, 26, 26, 10, 11, 12, 12, 9, 12, 12, 35], freeze="D2")
dv1 = DataValidation(type="list", formula1='"Must,Should,Could"', allow_blank=True); ws.add_data_validation(dv1); dv1.add(f"I2:I400")
dv2 = DataValidation(type="list", formula1='"Native,Config,Custom,Gap,TBD"', allow_blank=True); ws.add_data_validation(dv2); dv2.add("J2:L400")
dv3 = DataValidation(type="list", formula1='"Hypothesis,Validated,Changed,Dropped"', allow_blank=True); ws.add_data_validation(dv3); dv3.add("O2:O400")
dv4 = DataValidation(type="list", formula1='"Y,N,TBD"', allow_blank=True); ws.add_data_validation(dv4); dv4.add("M2:M400")

# ---- Observations
ws = wb.create_sheet("2 Observations")
sheet(ws, ["Obs ID", "Date", "Site", "Session", "Observation (what we saw / heard)", "Type", "Impact", "Related Req ID", "Source person / role", "Follow-up needed", "Photo / artifact ref"],
      [("OBS-001", "", "", "", "", "", "", "", "", "", "")], [9, 11, 11, 9, 70, 14, 10, 14, 20, 28, 18])
dv = DataValidation(type="list", formula1='"Pain point,Strength,Gap,Opportunity,Risk,Question"', allow_blank=True); ws.add_data_validation(dv); dv.add("F2:F500")
dv = DataValidation(type="list", formula1='"High,Medium,Low"', allow_blank=True); ws.add_data_validation(dv); dv.add("G2:G500")
dv = DataValidation(type="list", formula1='"SLO,Irvine"', allow_blank=True); ws.add_data_validation(dv); dv.add("C2:C500")

# ---- Irvine backlog
ws = wb.create_sheet("3 Irvine Backlog")
seed = [
 ("IRV-001", "Allocation", "Lines not fully committed / Special Order lines require manual reallocation before pick", "", "", "", "", "", "", "Validate frequency"),
 ("IRV-002", "Pick", "Partial pick closes task → reverse pick required to re-pick when stock arrives", "", "", "", "", "", "", "Validate"),
 ("IRV-003", "Issue", "Work Order Issue created manually; default inventory detail bin errors", "", "", "", "", "", "", "Who owns this step – QC vs. material handler?"),
 ("IRV-004", "Production", "No control preventing a run from starting without sufficient material or standard cost", "", "", "", "", "", "", "From Aug assessment"),
 ("IRV-005", "Timekeeping", "Timekeeping gaps in Manufacturing Mobile vs. Ezurio requirements", "", "", "", "", "", "", "Capture the requirement list"),
 ("IRV-006", "Quality", "RMA / rework / scrap tracked manually; QMS in SharePoint", "", "", "", "", "", "", ""),
 ("IRV-007", "De-kit", "Post-production reconciliation via templates + imports (IA / bin transfer / inventory transfer)", "", "", "", "", "", "", "Quantify effort"),
 ("IRV-008", "Automation", "Auto-Mark WOs Built workflow suspended 10/1 (error: BOM not defined for assembly item, WO 2447441)", "", "", "", "", "", "", "Why suspended? How are WOs closed now?"),
]
sheet(ws, ["ID", "Process step", "Issue / pain point", "Proposed optimization", "Type (Config / Report / Custom / Process)", "Value (H/M/L)", "Effort (S/M/L)", "Owner (Crowe / Y-Hire / Ezurio)", "ADO work item #", "Notes"],
      seed, [9, 14, 55, 40, 18, 10, 10, 18, 14, 30])
ws.freeze_panes = "C2"

# ---- GWTrack features
ws = wb.create_sheet("4 GWTrack Features")
feats = ["Job / work order creation & release", "Routing / operation sequencing", "Work instructions / document display", "Station scanning / operator transactions", "Material availability check", "Kitting / component issue", "Backflush / consumption",
         "Lot / serial genealogy", "Test data capture", "Defect capture & coded reasons", "Defect heatmaps", "NCR / rework / scrap", "RMA handling", "CAPA / 8D", "Labor / time capture", "WIP tracking & visibility",
         "Inventory transactions to NetSuite", "Cost / variance data", "Label printing", "Equipment / test-station integration", "Dashboards / reports", "User roles & permissions", "Audit trail / e-records", "Arena / BOM-ECO interface"]
sheet(ws, ["#", "GWTrack capability", "Demo'd? (Y/N)", "How it works today (brief)", "Automated vs. manual", "Used by (role)", "Frequency of use", "Business-critical? (Y/N)", "Where should it live? (GWTrack / NetSuite / Arena / MES / Retire)", "Gap / risk in alternatives", "Req ID"],
      [(i, f, "", "", "", "", "", "", "", "", "") for i, f in enumerate(feats, 1)], [5, 36, 10, 40, 16, 16, 12, 12, 24, 36, 10], freeze="C2")

# ---- Option scoring
ws = wb.create_sheet("5 Option Scoring")
ws.append(["Criteria weights sum to 100. Score each option 1 (poor) – 5 (excellent) after the onsite. Weighted score = SUMPRODUCT(weight, score)/100."])
ws.append([])
ws.append(["Criterion", "Weight", "A: Extend GWTrack", "B: Mfg Mobile / Adv Mfg", "C: 3rd-party MES / SuiteApp", "Basis / evidence"])
crit = [("Work execution & traceability fit", 15), ("Material availability / process control", 10), ("Standard cost enforcement", 5), ("Timekeeping / labor", 10), ("Quality tracking (defects, NCR, RMA, CAPA)", 15),
        ("ITAR / CUI & compliance fit", 15), ("Automation & performance retained", 10), ("Integration with NetSuite / Arena", 10), ("Scalability / supportability", 5), ("Estimated effort & risk of implementation", 5)]
for c, w in crit: ws.append([c, w, None, None, None, ""])
n = 3 + len(crit)
ws.append(["TOTAL", f"=SUM(B4:B{n})", f"=SUMPRODUCT($B$4:$B${n},C4:C{n})/100", f"=SUMPRODUCT($B$4:$B${n},D4:D{n})/100", f"=SUMPRODUCT($B$4:$B${n},E4:E{n})/100", ""])
for c in ws[3]: c.fill = HF; c.font = HFont
for row in ws.iter_rows(min_row=4):
    for c in row: c.font = BODY; c.border = BORDER; c.alignment = Alignment(wrap_text=True, vertical="top")
for c in ws[n + 1]: c.font = Font(name="Arial", size=9, bold=True)
ws["A1"].font = Font(name="Arial", size=9, italic=True)
for i, w in enumerate([44, 9, 18, 22, 24, 50], 1): ws.column_dimensions[get_column_letter(i)].width = w
dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True); ws.add_data_validation(dv); dv.add(f"C4:E{n}")

# ---- Open items
ws = wb.create_sheet("6 Open Items")
open_items = [
 ("OI-001", "Confirm ITAR/CUI scope and what may be shared with non-US Crowe personnel (incl. Y-Hire / international staff)", "Ezurio compliance / David England", "Before 10/19", "Open"),
 ("OI-002", "Confirm 'Jason' – role and written MES objectives", "David England", "Before 10/19", "Open"),
 ("OI-003", "Confirm attendee names/roles for each session (agenda placeholders)", "David England", "Before 10/16", "Open"),
 ("OI-004", "Pull current Irvine backlog from Azure DevOps (needs ADO org/project/PAT)", "Jordan Piatek", "Before 10/19", "Open"),
 ("OI-005", "Confirm Y-Hire (Hana, Kyle) availability for Wed 10/21 AM given India working hours", "Jordan Piatek", "Before 10/16", "Open"),
 ("OI-006", "Coverage for Tue 10/20 (Grant PTO) – Rohit leading technical sessions", "Jordan / Grant", "Before 10/19", "Open"),
 ("OI-007", "Plan for draft distribution while Jordan is on PTO Oct 23 – Nov 2", "Jordan / Grant / Sohaa", "10/21", "Open"),
 ("OI-008", "Data request sent to Ezurio (last page of the client agenda)", "Jordan Piatek", "10/9", "Open"),
 ("OI-009", "Confirm why Auto-Mark WOs Built workflow was suspended 10/1 and current WO closing process", "Trevor Kanwal / Ezurio", "Before 10/21", "Open"),
]
sheet(ws, ["ID", "Item", "Owner", "Due", "Status"], open_items, [9, 90, 32, 14, 10])
dv = DataValidation(type="list", formula1='"Open,In progress,Done,Blocked"', allow_blank=True); ws.add_data_validation(dv); dv.add("E2:E300")
wb.save("../03_Ezurio_Requirements_and_Observations_Tracker.xlsx"); print("saved")
