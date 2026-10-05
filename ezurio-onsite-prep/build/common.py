"""Shared helpers + session content for the Ezurio onsite prep materials."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_COLOR_INDEX
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import re

NAVY = "002060"; BLUE = "1F4E79"; LBLUE = "2E75B6"; GREY = "F2F2F2"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def set_cell_width(cell, inches):
    cell.width = Inches(inches)
    tcPr = cell._tc.get_or_add_tcPr()
    w = tcPr.find(qn("w:tcW"))
    if w is None:
        w = OxmlElement("w:tcW"); tcPr.append(w)
    w.set(qn("w:w"), str(int(inches * 1440))); w.set(qn("w:type"), "dxa")

def new_doc(margin=0.75):
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Inches(margin)
    st = d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    st.paragraph_format.space_after = Pt(4)
    for name, size, color in (("Heading 1", 16, BLUE), ("Heading 2", 13, BLUE), ("Heading 3", 11, LBLUE)):
        h = d.styles[name]; h.font.name = "Arial"; h.font.size = Pt(size); h.font.bold = True
        h.font.color.rgb = RGBColor.from_string(color)
        h.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial"); h.element.rPr.rFonts.set(qn("w:ascii"), "Arial"); h.element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        h.paragraph_format.space_before = Pt(10); h.paragraph_format.space_after = Pt(4)
    for ln in ("List Bullet", "List Bullet 2"):
        d.styles[ln].font.name = "Arial"; d.styles[ln].font.size = Pt(10)
        d.styles[ln].paragraph_format.space_after = Pt(2)
    return d

def add_rich(p, text, size=None, italic=False, bold=False, color=None):
    """Text with [[placeholder]] segments highlighted yellow."""
    for part in re.split(r"(\[\[.*?\]\])", text):
        if not part: continue
        ph = part.startswith("[[")
        r = p.add_run("[" + part[2:-2] + "]" if ph else part)
        r.italic = italic; r.bold = bold
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = RGBColor.from_string(color)
        if ph: r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p

def para(d, text="", size=None, italic=False, bold=False, color=None, style=None, after=None):
    p = d.add_paragraph(style=style)
    add_rich(p, text, size, italic, bold, color)
    if after is not None: p.paragraph_format.space_after = Pt(after)
    return p

def bullet(d, text, level=1, italic=False):
    p = d.add_paragraph(style="List Bullet" if level == 1 else "List Bullet 2")
    add_rich(p, text, italic=italic)
    return p

def banner(d, text, sub=None, width=7.0):
    t = d.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]; shade(c, BLUE); set_cell_width(c, width)
    p = c.paragraphs[0]; r = p.add_run(text); r.bold = True; r.font.size = Pt(12); r.font.color.rgb = RGBColor(255, 255, 255)
    if sub:
        p2 = c.add_paragraph(); r2 = p2.add_run(sub); r2.font.size = Pt(9); r2.font.color.rgb = RGBColor(255, 255, 255)
    d.add_paragraph().paragraph_format.space_after = Pt(2)

def grid_table(d, headers, rows, widths, header_fill=NAVY, size=9):
    t = d.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; shade(c, header_fill); set_cell_width(c, widths[i])
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = RGBColor(255, 255, 255)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            set_cell_width(cells[i], widths[i])
            lines = val if isinstance(val, list) else [val]
            first = True
            for ln in lines:
                p = cells[i].paragraphs[0] if first else cells[i].add_paragraph()
                first = False
                p.paragraph_format.space_after = Pt(1)
                add_rich(p, ln, size=size, bold=(i == 1 and not isinstance(val, list)))
        if row[1] and str(row[1]).lower().startswith(("break", "lunch")):
            for c in cells: shade(c, GREY)
    return t

# ---------------------------------------------------------------- schedule
# (time, session id, session title, ezurio attendees, [client-facing topics])
DAY1 = [
 ("8:00 – 8:45", "1.1", "Kickoff & Objectives", "[[David England, John Stover, Gateworks site leadership, TBC]]",
  ["Goals for the three days and how the output feeds the next phase", "Confirm participants and who owns which area at Gateworks",
   "What does \"keeping today's level of automation and performance\" mean to the Gateworks team?"]),
 ("8:45 – 10:45", "1.2", "Gateworks Operations Walk", "[[Operations / production lead, quality lead, supervisors]]",
  ["Floor walk from receiving through build, test and ship", "How work is released, tracked and completed on the floor today",
   "Where operators and supervisors touch GWTrack (stations, scanners, screens)", "Pain points and workarounds the team lives with"]),
 ("10:45 – 11:00", "", "Break / Buffer", "", []),
 ("11:00 – 12:30", "1.3", "GWTrack Demo I – Production Execution & Traceability", "[[GWTrack power users / admin, production control]]",
  ["Live demo: work orders, routings, work centers and shop-floor execution", "Material availability checks and kitting / issue to the floor",
   "Traceability and genealogy (lot / serial / as-built records)", "Dashboards and reports used day to day"]),
 ("12:30 – 1:15", "", "Lunch (catered in)", "", []),
 ("1:15 – 3:15", "1.4", "GWTrack Demo II – Quality, Labor & Inventory", "[[Quality lead, supervisors, inventory / materials]]",
  ["Quality tracking: defect capture, defect heatmaps, NCR / rework / scrap", "Customer returns (RMA) and corrective action",
   "Labor and time capture – how hours are recorded and used", "Inventory, stock locations and costing touchpoints with NetSuite"]),
 ("3:15 – 3:30", "", "Break / Buffer", "", []),
 ("3:30 – 4:30", "1.5", "Day 1 Wrap-up & Open Items", "[[David England, John Stover, site leadership]]",
  ["Recap what we saw and what we still need to see", "Confirm Day 2 attendees and any documents to pull"]),
]
DAY2 = [
 ("8:00 – 10:00", "2.1", "Systems Landscape & Integrations", "[[IT / NetSuite admin, GWTrack admin, engineering (Arena)]]",
  ["How GWTrack, NetSuite and Arena work together today – what lives where", "Data flows and hand-offs: BOMs / engineering changes, items, orders, inventory, costs",
   "Integrations, scripts, scheduled jobs and manual bridges (spreadsheets, re-keying)", "Reports and analytics the business relies on"]),
 ("10:00 – 10:15", "", "Break / Buffer", "", []),
 ("10:15 – 12:00", "2.2", "Compliance, Security & Data Requirements", "[[Compliance / contracts lead, IT security]]",
  ["Customer and contract requirements that shape system choices (ITAR / CUI)", "Access control, user roles, hosting and data handling expectations",
   "Audit and record-keeping needs on the shop floor", "Who needs to be involved in validating compliance questions"]),
 ("12:00 – 12:45", "", "Lunch (catered in)", "", []),
 ("12:45 – 2:45", "2.3", "Capabilities to Keep: Automation & Performance", "[[David England, John Stover, site leadership, power users]]",
  ["Walk through the \"must not lose\" capabilities and why they matter", "Where the current system is fast / automated vs. manual",
   "Manufacturing system objectives for the next phase and what success looks like", "Scalability: what would need to change to support additional volume, users or sites"]),
 ("2:45 – 3:00", "", "Break / Buffer", "", []),
 ("3:00 – 4:30", "2.4", "Requirements Read-back & Prioritization", "[[David England, John Stover, site leadership]]",
  ["Read back requirements captured over Days 1–2", "Prioritize: must have / should have / nice to have", "Open questions, follow-ups and owners"]),
]
DAY3 = [
 ("8:00 – 8:30", "3.1", "Kickoff & Goals for Irvine", "[[David England, John Stover, Irvine operations leadership]]",
  ["Goals for the day and how findings feed the improvement backlog", "What has changed since the July 2025 go-live"]),
 ("8:30 – 10:00", "3.2", "Floor Walk & Live Process Observation", "[[Warehouse operators and manager, QC, production operators]]",
  ["Observe a work order flow from allocation, pick and staging through issue and production", "Warehouse (WMS) and Manufacturing Mobile in use at the point of work",
   "Transfers to production input bin, lab input and back to warehouse"]),
 ("10:00 – 10:15", "", "Break / Buffer", "", []),
 ("10:15 – 12:00", "3.3", "Manufacturing Mobile: Process Controls & Timekeeping", "[[Production manager, planner, cost accounting, QC lead]] + Crowe NetSuite support team (remote)",
  ["Material availability checks before a run starts", "Standard cost setup and enforcement", "Operator time capture: badge-in, pause / end, completions, exceptions",
   "What the team needs from timekeeping (labor reporting, costing, payroll hand-off)"]),
 ("12:00 – 12:45", "", "Lunch (catered in)", "", []),
 ("12:45 – 2:00", "3.4", "Quality, Rework / RMA / Scrap & Post-Production Reconciliation", "[[QC lead, warehouse manager, production manager]]",
  ["Quality tracking in SharePoint QMS and manual RMA / rework / scrap logs", "De-kit and post-production reconciliation of issued components",
   "Reporting the team wants but cannot get today"]),
 ("2:00 – 2:15", "", "Break / Buffer", "", []),
 ("2:15 – 3:45", "3.5", "Optimization Workshop & Backlog", "[[Irvine leads]] + Crowe NetSuite support team (as available)",
  ["Pain points and improvement ideas surfaced today", "Quick wins vs. larger changes; effort and value", "Draft backlog and owners"]),
 ("3:45 – 4:30", "3.6", "Week Wrap-up & Next Steps", "[[David England, John Stover]]",
  ["Recap of observations across Gateworks and Irvine", "Open items, documents to follow up on, and timing of the summary"]),
]
DAYS = [("DAY 1 – Gateworks (San Luis Obispo): GWTrack Demo & Shop-Floor Review", "Monday, October 19", "3026 South Higuera Street, San Luis Obispo, CA 93401", DAY1),
        ("DAY 2 – Gateworks (San Luis Obispo): Systems, Compliance & Requirements", "Tuesday, October 20", "3026 South Higuera Street, San Luis Obispo, CA 93401", DAY2),
        ("DAY 3 – Irvine: Process Optimization", "Wednesday, October 21", "16450 Scientific, Irvine, CA 92618", DAY3)]
